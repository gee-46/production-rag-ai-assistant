"""
Text extraction per document type.

The old code only handled .txt and .docx, read the whole upload into memory
with `await file.read()`, and wrote docx uploads to a shared hardcoded
`temp.docx` path in the process's working directory — a race condition
under concurrent uploads. Extraction here always operates on a file already
saved to a unique per-document path (see ingestion/pipeline.py) and adds PDF
support with an OCR fallback for scanned/image-only pages.
"""
from __future__ import annotations

from pathlib import Path

from app.core.errors import ValidationAppError
from app.core.logging import get_logger

logger = get_logger(__name__)

MIN_CHARS_PER_PAGE_BEFORE_OCR = 20  # below this, assume the PDF page is a scanned image


def extract_text(path: Path, content_type: str) -> str:
    suffix = path.suffix.lower()
    if suffix in (".txt", ".md"):
        return _extract_plain_text(path)
    if suffix == ".docx":
        return _extract_docx(path)
    if suffix == ".pdf":
        return _extract_pdf(path)
    raise ValidationAppError(f"Unsupported file type: {suffix}")


def _extract_plain_text(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="replace")


def _extract_docx(path: Path) -> str:
    from docx import Document as DocxDocument

    doc = DocxDocument(str(path))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    # Tables are common in real-world docx files and were silently dropped
    # before; include their cell text too.
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                paragraphs.append(row_text)
    return "\n\n".join(paragraphs)


def _extract_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages_text: list[str] = []
    ocr_pages = 0

    for i, page in enumerate(reader.pages):
        text = (page.extract_text() or "").strip()
        if len(text) < MIN_CHARS_PER_PAGE_BEFORE_OCR:
            ocr_text = _ocr_page(path, i)
            if ocr_text:
                text = ocr_text
                ocr_pages += 1
        pages_text.append(text)

    if ocr_pages:
        logger.info("pdf_ocr_fallback_used path=%s pages=%s", path.name, ocr_pages)

    return "\n\n".join(p for p in pages_text if p)


def _ocr_page(path: Path, page_index: int) -> str:
    """
    OCR a single PDF page via pdf2image + pytesseract. Requires system
    binaries (poppler-utils, tesseract-ocr) that are installed in the
    Docker image but may not be present in every dev environment — OCR is
    attempted best-effort and silently skipped (with a warning) rather than
    failing the whole ingestion job if those binaries are missing.
    """
    try:
        import pytesseract
        from pdf2image import convert_from_path

        images = convert_from_path(str(path), first_page=page_index + 1, last_page=page_index + 1)
        if not images:
            return ""
        return pytesseract.image_to_string(images[0])
    except Exception as exc:  # noqa: BLE001
        logger.warning("ocr_unavailable_or_failed path=%s page=%s error=%s", path.name, page_index, exc)
        return ""
